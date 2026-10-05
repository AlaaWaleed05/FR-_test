"""Render docs/bank-hosting-requirements.md onto the AZ Sudan letterhead."""
import re
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MD = sys.argv[1]
TEMPLATE = sys.argv[2]
OUT = sys.argv[3]
DIAGRAM = sys.argv[4]

PURPLE = RGBColor(0x42, 0x27, 0x74)
RED = RGBColor(0xE4, 0x31, 0x2A)
INK = RGBColor(0x20, 0x22, 0x26)
GREY = RGBColor(0x5A, 0x5E, 0x66)
PURPLE_HEX = "422774"
BAND_HEX = "EFEDF4"
ZEBRA_HEX = "F7F7F9"

doc = Document(TEMPLATE)

# ---------------------------------------------------------------- page set-up
# The letterhead image is anchored to the MARGIN, so moving a margin would move the
# artwork with it. Re-anchor to the page first, then the margins are free.
hdr = doc.sections[0].header.part.element
for tag, off in (("positionH", "0"), ("positionV", "0")):
    for node in hdr.iter(qn("wp:" + tag)):
        node.set("relativeFrom", "page")
        node.find(qn("wp:posOffset")).text = off

s = doc.sections[0]
s.top_margin = Cm(4.6)      # clears the header artwork
s.bottom_margin = Cm(2.3)   # clears the footer band
s.left_margin = Cm(1.9)
s.right_margin = Cm(1.9)
CONTENT_W = Cm(21) - s.left_margin - s.right_margin


def style(name, size, bold=False, colour=INK, before=0, after=6, keep=False, italic=False):
    st = doc.styles.add_style(name, 1)
    st.base_style = doc.styles["Normal"]
    f = st.font
    f.name = "Calibri"
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:cs"), "Calibri")
    rf.set(qn("w:eastAsia"), "Calibri")
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = colour
    p = st.paragraph_format
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.keep_with_next = keep
    p.widow_control = True
    return st


style("DocTitle", 26, True, PURPLE, 0, 4)
style("DocSub", 13, False, GREY, 0, 18)
style("H1", 15, True, PURPLE, 18, 8, keep=True)
style("H2", 11.5, True, PURPLE, 12, 5, keep=True)
style("Body", 10, False, INK, 0, 7)
style("Bullet", 10, False, INK, 0, 4)
style("Mono", 9, False, INK, 4, 8)
style("Caption", 8.5, False, GREY, 4, 12, italic=True)
style("Cell", 9, False, INK, 1, 1)
style("CellHead", 9, True, RGBColor(0xFF, 0xFF, 0xFF), 1, 1)
doc.styles["Mono"].font.name = "Consolas"


def shade(el, hexcolour):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolour)
    el.append(sh)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "C6C9D0")
        b.append(e)
    tblPr.append(b)


BOLD = re.compile(r"(\*\*.+?\*\*)")
CODE = re.compile(r"(`[^`]+`)")


def _emit(par, text, bold, base_size, base_colour):
    """Code spans nest inside bold, so bold is carried down rather than split on."""
    for piece in CODE.split(text):
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt((base_size or 10) - 0.5)
        else:
            r = par.add_run(piece)
            if base_size:
                r.font.size = Pt(base_size)
        r.bold = bold
        if base_colour:
            r.font.color.rgb = base_colour


def runs(par, text, base_size=None, base_colour=None):
    for piece in BOLD.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            _emit(par, piece[2:-2], True, base_size, base_colour)
        else:
            _emit(par, piece, False, base_size, base_colour)


def para(text, stylename="Body", **kw):
    p = doc.add_paragraph(style=stylename)
    runs(p, text, **kw)
    return p


def bullet(text, numbered=False):
    p = doc.add_paragraph(style="Bullet")
    p.paragraph_format.left_indent = Cm(1.1)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    m = re.match(r"^(\d+)\. (.*)$", text)
    if m:
        lead, text = m.group(1) + ".  ", m.group(2)
    else:
        lead = "•  "
    r = p.add_run(lead)
    r.font.color.rgb = PURPLE
    r.bold = True
    runs(p, text)
    return p


def add_table(rows, has_header):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.autofit = True
    borders(t)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        head = has_header and i == 0
        for j, text in enumerate(row):
            cell = cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.style = doc.styles["CellHead" if head else "Cell"]
            if not has_header and j == 0:
                text = "**" + text + "**" if text and "**" not in text else text
            runs(p, text, base_size=9,
                 base_colour=RGBColor(0xFF, 0xFF, 0xFF) if head else None)
            if head:
                shade(cell._tc.get_or_add_tcPr(), PURPLE_HEX)
            elif i % 2 == 0:
                shade(cell._tc.get_or_add_tcPr(), ZEBRA_HEX)
        if head:
            trPr = cells[0]._tc.getparent().get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
        trPr = cells[0]._tc.getparent().get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    doc.add_paragraph(style="Body").paragraph_format.space_after = Pt(2)
    return t


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


# ---------------------------------------------------------------- read source
raw = open(MD, encoding="utf8").read().splitlines()

# A wrapped list item is one item, not an item followed by a paragraph. Fold the
# indented continuation lines back into the line they belong to, leaving fenced
# blocks untouched.
src = []
in_fence = False
last_was_list = False
for ln in raw:
    if ln.strip().startswith("```"):
        in_fence = not in_fence
        src.append(ln)
        last_was_list = False
        continue
    if in_fence:
        src.append(ln)
        continue
    body = ln.strip()
    is_list = body.startswith("- ") or bool(re.match(r"^\d+\. ", body))
    if (last_was_list and ln.startswith("  ") and body
            and not is_list and not body.startswith("|")):
        src[-1] = src[-1].rstrip() + " " + body
        continue
    src.append(ln)
    last_was_list = is_list

# cover
title = "Bayanati"
subtitle = "Hosting requirements for a bank-operated deployment"

p = doc.add_paragraph(style="DocTitle")
p.add_run(title)
para(subtitle, "DocSub")

i = 0
while not src[i].startswith("| | |"):
    i += 1
meta = []
i += 2  # skip header separator row
while src[i].startswith("|"):
    meta.append(split_row(src[i]))
    i += 1
add_table([[c.replace("**", "") for c in r] for r in meta], has_header=False)

# skip to first "## "
while not src[i].startswith("## "):
    i += 1

buf = []
pending_table = []


def flush_para():
    global buf
    if buf:
        text = " ".join(buf).strip()
        if text:
            para(text)
        buf = []


def flush_table():
    global pending_table
    if pending_table:
        header = pending_table[0]
        body = pending_table[2:]
        has_header = any(c for c in header)
        add_table(([header] if has_header else []) + body, has_header)
        pending_table = []


n = len(src)
while i < n:
    line = src[i]
    stripped = line.strip()

    if stripped.startswith("```"):
        flush_para()
        flush_table()
        fence = []
        i += 1
        while i < n and not src[i].strip().startswith("```"):
            fence.append(src[i])
            i += 1
        i += 1
        block = "\n".join(fence)
        if "PUBLIC INTERNET" in block:
            doc.add_picture(DIAGRAM, width=CONTENT_W)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para("Figure 1. Production topology. Three servers, the zones they sit in, "
                 "and every connection in and out.", "Caption")
        else:
            for ln in fence:
                pm = doc.add_paragraph(style="Mono")
                pm.paragraph_format.left_indent = Cm(0.5)
                pm.paragraph_format.space_after = Pt(0)
                pm.add_run(ln)
            doc.add_paragraph(style="Body").paragraph_format.space_after = Pt(2)
        continue

    if stripped.startswith("|"):
        flush_para()
        pending_table.append(split_row(stripped))
        i += 1
        continue
    flush_table()

    if stripped.startswith("### "):
        flush_para()
        para(stripped[4:], "H2")
    elif stripped.startswith("## "):
        flush_para()
        para(stripped[3:], "H1")
    elif stripped.startswith("- "):
        flush_para()
        bullet(stripped[2:])
    elif re.match(r"^\d+\. ", stripped):
        flush_para()
        bullet(stripped)
    elif stripped == "---" or stripped == "":
        flush_para()
    else:
        buf.append(stripped)
    i += 1

flush_para()
flush_table()

doc.save(OUT)
print("saved", OUT)
